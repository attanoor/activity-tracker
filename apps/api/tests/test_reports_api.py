import io
from datetime import datetime, time
from zoneinfo import ZoneInfo

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from app.ai import llm as llm_mod
from app.main import app
from app.models import AuditLog, CalendarEvent, DsrReport, Team, WorkAction
from conftest import login
from test_dsr import DAY, FakeLLM, person  # noqa: F401  (fixture)


@pytest.fixture(autouse=True)
def fake_llm():
    from app.dsr import generation
    llm_mod.override, generation.inline = FakeLLM(), True
    yield
    llm_mod.override, generation.inline = None, False


def test_generate_edit_submit(client, person):  # noqa: F811
    login(client, "ali@example.com")
    assert client.post(f"/me/dsr/{DAY}/generate").json() == {"generating": False}
    o = client.get(f"/me/overview/{DAY}").json()
    assert o["generating"] is False and o["generation_error"] is None
    d = o["dsr"]
    assert d["status"] == "draft" and d["content"]["tasks"]
    task = d["content"]["tasks"][0]
    secs = task["seconds"]
    r = client.put(f"/me/dsr/{DAY}", json={"tasks": [{"id": task["id"], "title": "My wording", "summary": "Mine",
                                                      "status": "done", "seconds": 1}],
                                         "highlights": ["Fixed login"], "blockers": [], "tomorrow": ["Deploy"],
                                         "insights": "Fine day."})
    assert r.status_code == 200 and r.json()["edited"]
    t = r.json()["content"]["tasks"][0]
    assert (t["title"], t["status"], t["seconds"]) == ("My wording", "done", secs)  # time can't be edited
    bad = client.put(f"/me/dsr/{DAY}", json={"tasks": [{"id": "t99", "title": "Invented", "status": "done"}]})
    assert bad.status_code == 422
    assert client.post(f"/me/dsr/{DAY}/submit").json()["status"] == "submitted"
    assert client.put(f"/me/dsr/{DAY}", json={"tasks": []}).status_code == 409
    assert client.post(f"/me/dsr/{DAY}/generate").status_code == 409


def test_overview_and_meeting_confirmation(client, person, db):  # noqa: F811
    login(client, "ali@example.com")
    o = client.get(f"/me/overview/{DAY}").json()
    assert sum(o["analysis"]["breakdown"].values()) == 9 * 3600
    [m] = o["analysis"]["meetings"]
    assert m["status"] == "attended"  # PC unused during the standup: a phone or room meeting
    assert o["analysis"]["breakdown"]["meeting"] == 900
    ev = db.query(CalendarEvent).one()
    assert client.post(f"/me/meetings/{ev.id}/attendance", json={"attended": False}).status_code == 200
    o = client.get(f"/me/overview/{DAY}").json()
    assert o["analysis"]["meetings"][0]["status"] == "not_attended" and o["analysis"]["breakdown"]["meeting"] == 0
    assert client.get(f"/me/trend?end={DAY}&days=3").json()["days"][-1]["breakdown"]["meeting"] == 0


def test_manager_access_hides_private_bits_and_is_audited(client, person, db, make_user):  # noqa: F811
    team = db.get(Team, person.team_id)
    boss = make_user("boss@example.com", role="manager")
    other = make_user("other@example.com", role="manager")
    team.manager_id = boss.id
    db.add(WorkAction(ext_key="gmail:1", user_id=person.id, day=DAY, source="gmail", action="replied", object_id="t1",
                      object_title="Salary question", at=datetime(2026, 10, 21, 6, tzinfo=ZoneInfo("UTC")), meta={}))
    db.commit()
    login(client, "ali@example.com")
    client.post(f"/me/dsr/{DAY}/generate")
    client.post("/auth/logout")

    login(client, "boss@example.com")
    o = client.get(f"/people/{person.id}/overview/{DAY}").json()
    assert all(a["title"] == "(email)" for a in o["actions"] if a["source"] == "gmail")
    assert "suggestions" not in o["dsr"]["content"]
    assert "Salary" not in str(o["dsr"]["content"])
    assert db.query(AuditLog).filter_by(action="view.day", actor_id=boss.id).count() == 1
    rows = client.get(f"/team/{team.id}/day/{DAY}").json()["rows"]
    assert [r["name"] for r in rows] == ["ali"] and rows[0]["dsr"] == "draft"
    s = client.post(f"/team/{team.id}/day/{DAY}/summary").json()
    assert s["content"]["summary"]
    client.post("/auth/logout")

    login(client, "other@example.com")
    assert client.get(f"/people/{person.id}/overview/{DAY}").status_code == 404
    assert client.get(f"/team/{team.id}/day/{DAY}").status_code == 404
    client.post("/auth/logout")
    make_user("peer@example.com")
    login(client, "peer@example.com")
    assert client.get(f"/people/{person.id}/overview/{DAY}").status_code == 404


def test_exports(client, person):  # noqa: F811
    login(client, "ali@example.com")
    client.post(f"/me/dsr/{DAY}/generate")
    x = client.get(f"/me/dsr/{DAY}/export/xlsx")
    wb = load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames == ["DSR", "Work Output", "Meetings", "Timeline", "9-Hour Breakdown", "By App & Site", "Hourly"]
    assert wb["9-Hour Breakdown"]._charts and wb["By App & Site"]._charts
    total = sum(r[1].value for r in wb["9-Hour Breakdown"].iter_rows(min_row=2))
    assert round(total, 1) == 9.0
    p = client.get(f"/me/dsr/{DAY}/export/pdf")
    assert p.content[:4] == b"%PDF"
    assert client.get(f"/me/dsr/{DAY}/export/template").status_code == 404  # none configured


def template_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "DSR"
    ws["A1"], ws["A2"] = "Name", "Date"
    ws["A5"], ws["B5"], ws["C5"] = "Task", "Hours", "Status"
    ws["B6"].font = Font(bold=True, color="FF0000")
    ws["E1"] = "=SUM(B6:B20)"
    wb.create_sheet("Notes")["A1"] = "keep me"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_company_template_flow(client, person, make_user):  # noqa: F811
    make_user("admin@example.com", role="admin")
    login(client, "ali@example.com")
    client.post(f"/me/dsr/{DAY}/generate")
    client.post("/auth/logout")
    login(client, "admin@example.com")
    t = client.post("/admin/templates", data={"name": "Company DSR"},
                    files={"file": ("dsr.xlsx", template_bytes(), "application/octet-stream")}).json()
    assert t["sheets"] == ["DSR", "Notes"]
    assert client.put(f"/admin/templates/{t['id']}/mapping", json={"cells": {"name": "Z"}}).status_code == 422
    mapping = {"sheet": "DSR", "cells": {"name": "B1", "date": "B2"}, "start_row": 6,
               "columns": {"task": "A", "hours": "B", "task_status": "C"}}
    assert client.put(f"/admin/templates/{t['id']}/mapping", json=mapping).status_code == 200
    rows = client.post(f"/admin/templates/{t['id']}/preview", json={"user_id": person.id, "day": str(DAY)}).json()["rows"]
    assert rows[0][1] == "ali" and rows[5][0]
    assert client.post(f"/admin/templates/{t['id']}/activate").status_code == 200
    client.post("/auth/logout")
    login(client, "ali@example.com")
    wb = load_workbook(io.BytesIO(client.get(f"/me/dsr/{DAY}/export/template").content))
    ws = wb["DSR"]
    assert ws["B1"].value == "ali" and ws["E1"].value == "=SUM(B6:B20)" and wb["Notes"]["A1"].value == "keep me"
    assert ws["B6"].font.bold and ws["A6"].value


def test_master_workbook_has_no_duplicates(db, person, tmp_path):  # noqa: F811
    from app.dsr.service import generate
    from app.exports.bundle import gather
    from app.exports.excel import append_master
    generate(db, person, DAY)
    path = tmp_path / "master.xlsx"
    n1 = append_master(path, gather(db, person, DAY))
    append_master(path, gather(db, person, DAY))
    ws = load_workbook(path)["2026-10"]
    assert ws.max_row == 1 + n1


def test_worker_drafts_after_shift_and_mails_manager(db, person, make_user, monkeypatch):  # noqa: F811
    from app import settings_store, worker
    team = db.get(Team, person.team_id)
    boss = make_user("boss@example.com", role="manager")
    team.manager_id = boss.id
    db.commit()
    settings_store.update(db, "email", {"enabled": True, "host": "smtp.local", "from_addr": "dsr@example.com",
                                        "send_time": "19:00"})
    sent = []
    monkeypatch.setattr(worker.mailer, "send", lambda db, to, subject, text, html, files=(), cc=(): sent.append((to, subject, [f[0] for f in files], cc)))
    now_local = {"t": datetime.combine(DAY, time(17, 0), ZoneInfo("Asia/Karachi"))}
    monkeypatch.setattr(worker, "local_now", lambda tz: now_local["t"].astimezone(ZoneInfo(tz)))
    monkeypatch.setattr(worker.runner, "sync_google", lambda *a, **k: {})
    monkeypatch.setattr(worker.runner, "sync_mantis", lambda *a, **k: {})
    from app.models import Device
    db.query(Device).update({"revoked_at": None})
    db.commit()

    worker.job_dsr(db, now_local["t"])
    assert db.query(DsrReport).count() == 0  # shift not over
    now_local["t"] = now_local["t"].replace(hour=18, minute=20)
    worker.job_dsr(db, now_local["t"])
    worker.job_dsr(db, now_local["t"])
    assert db.query(DsrReport).count() == 1
    now_local["t"] = now_local["t"].replace(hour=19, minute=5)
    worker.job_mail(db, now_local["t"])
    worker.job_mail(db, now_local["t"])
    dsr_mails = [m for m in sent if m[1].startswith("DSR")]
    assert len(dsr_mails) == 1
    to, subject, files, cc = dsr_mails[0]
    assert to == ["boss@example.com"] and cc == ["ali@example.com"] and "not reviewed" in subject
    assert any(f.endswith(".xlsx") for f in files) and any(f.endswith(".pdf") for f in files)
    assert [m for m in sent if m[1].startswith("Team")]
    assert db.query(DsrReport).one().status == "sent"
