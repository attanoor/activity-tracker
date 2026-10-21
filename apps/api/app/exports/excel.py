"""Excel output: the tracker's own DSR workbook, the company template filled in, and the master workbook."""

import io
from copy import copy
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
from openpyxl.utils.exceptions import CellCoordinatesException

from app.analysis.engine import BUCKETS
from app.exports.bundle import Bundle, day_fields, hours, local, task_fields

HEAD = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="12805A")
WRAP = Alignment(wrap_text=True, vertical="top")
BUCKET_LABELS = {"productive": "Productive", "meeting": "Meetings/Calls", "communication": "Email & chat",
                 "neutral": "Neutral/Admin", "non_work": "Non-work", "idle": "Idle at desk", "private": "Private/paused",
                 "break": "Break", "away": "Away/untracked"}


def _table(ws, header: list[str], rows: list[list], widths: list[int] | None = None) -> None:
    ws.append(header)
    for c in ws[ws.max_row]:
        c.font, c.fill = HEAD, HEAD_FILL
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths or [], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=ws.max_row - len(rows) + 1, column=1) if rows else None


def dsr_workbook(bundles: list[Bundle], team_name: str | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "DSR"
    rows = []
    for b in bundles:
        f = day_fields(b)
        for t in b.content.get("tasks", []):
            tf = task_fields(t)
            rows.append([b.day, b.user.name, tf["task"], tf["task_summary"], tf["hours"], tf["task_status"],
                         tf["category"], tf["tickets"]])
        rows.append([b.day, b.user.name, "Highlights", f["highlights"], None, None, None, None])
        rows.append([b.day, b.user.name, "Blockers", f["blockers"] or "None", None, None, None, None])
        rows.append([b.day, b.user.name, "Tomorrow", f["tomorrow"], None, None, None, None])
        rows.append([b.day, b.user.name, "Insights (AI)", f["ai_summary"], None, None, None, None])
        rows.append([b.day, b.user.name, "Suggestions (AI)", f["ai_suggestions"], None, None, None, None])
    _table(ws, ["Date", "Person", "Task", "Summary", "Hours", "Status", "Category", "Tickets"], rows,
           [12, 20, 40, 70, 8, 12, 18, 14])
    for row in ws.iter_rows(min_row=2):
        row[3].alignment = WRAP

    ws = wb.create_sheet("Work Output")
    counts, actions = [], []
    for b in bundles:
        f = day_fields(b)
        counts.append([b.day, b.user.name, f["tickets_reported"], f["tickets_responded"], f["tickets_resolved"],
                       f["tickets_touched"], f["emails_sent"], f["email_replies"], f["email_threads"],
                       f["docs_created"], f["docs_edited"], f["meeting_count"], f["meeting_hours"]])
        for a in b.actions:
            actions.append([b.day, b.user.name, local(a.at, b.tz), a.source, a.action, a.object_id, a.object_title])
    _table(ws, ["Date", "Person", "Tickets reported", "Ticket responses", "Tickets resolved", "Tickets touched",
                "Emails sent", "Email replies", "Email threads", "Docs created", "Docs edited", "Meetings",
                "Meeting hours"], counts, [12, 20] + [11] * 11)
    ws.append([])
    ws.append(["Every action"])
    ws[ws.max_row][0].font = Font(bold=True)
    ws.append(["Date", "Person", "Time", "Source", "Action", "Object", "Title"])
    for c in ws[ws.max_row]:
        c.font, c.fill = HEAD, HEAD_FILL
    for r in actions:
        ws.append(r)

    ws = wb.create_sheet("Meetings")
    _table(ws, ["Date", "Person", "Start", "End", "Title", "Minutes", "Status"],
           [[b.day, b.user.name, m["start"][11:16], m["end"][11:16], m["title"], round(m["duration_s"] / 60), m["status"]]
            for b in bundles for m in b.analysis["meetings"]], [12, 20, 8, 8, 40, 9, 14])

    ws = wb.create_sheet("Timeline")
    _table(ws, ["Date", "Person", "Start", "End", "Minutes", "Class", "Category", "App", "Site", "Title"],
           [[b.day, b.user.name, local(s.started_at, b.tz), local(s.ended_at, b.tz), round(s.duration_s / 60, 1),
             s.klass, getattr(s, "category", None), s.app, s.domain, s.title] for b in bundles for s in b.sessions],
           [12, 20, 8, 8, 9, 14, 18, 18, 24, 60])

    ws = wb.create_sheet("9-Hour Breakdown")
    _table(ws, ["Bucket"] + [b.user.name for b in bundles],
           [[BUCKET_LABELS[k]] + [hours(b.analysis["breakdown"][k]) for b in bundles] for k in BUCKETS],
           [20] + [14] * len(bundles))
    if bundles:
        pie = PieChart()
        pie.title = f"Where the shift went: {bundles[0].user.name} (hours)"
        pie.add_data(Reference(ws, min_col=2, min_row=1, max_row=len(BUCKETS) + 1), titles_from_data=True)
        pie.set_categories(Reference(ws, min_col=1, min_row=2, max_row=len(BUCKETS) + 1))
        pie.height, pie.width = 9, 14
        ws.add_chart(pie, f"{get_column_letter(len(bundles) + 3)}2")

    ws = wb.create_sheet("By App & Site")
    tops = []
    for b in bundles:
        for t in b.analysis["top_targets"]:
            tops.append([b.day, b.user.name, t["target"], t.get("category") or "", t["klass"], round(t["seconds"] / 60)])
    _table(ws, ["Date", "Person", "App or site", "Category", "Class", "Minutes"], tops, [12, 20, 30, 18, 14, 10])
    if tops:
        bar = BarChart()
        bar.type, bar.title, bar.y_axis.title = "bar", "Minutes by app or site", "Minutes"
        n = min(len(tops), 10)
        bar.add_data(Reference(ws, min_col=6, min_row=1, max_row=n + 1), titles_from_data=True)
        bar.set_categories(Reference(ws, min_col=3, min_row=2, max_row=n + 1))
        bar.height, bar.width, bar.legend = 9, 16, None
        ws.add_chart(bar, "H2")

    ws = wb.create_sheet("Hourly")
    hourly = []
    for b in bundles:
        for h in b.analysis["hours"]:
            hourly.append([b.day, b.user.name, f"{h['hour']:02d}:00"] + [round(h[k] / 60) for k in BUCKETS])
    _table(ws, ["Date", "Person", "Hour"] + [BUCKET_LABELS[k] for k in BUCKETS], hourly, [12, 20, 8] + [13] * len(BUCKETS))
    if hourly:
        last = ws.max_row
        for col in range(4, 4 + len(BUCKETS)):
            letter = get_column_letter(col)
            ws.conditional_formatting.add(f"{letter}2:{letter}{last}", ColorScaleRule(
                start_type="num", start_value=0, start_color="FFFFFF", end_type="num", end_value=60, end_color="12805A"))

    if team_name is not None:
        ws = wb.create_sheet("Team Summary", 0)
        _table(ws, ["Person", "Report", "Worked h", "Productive h", "Meetings h", "Email/chat h", "Idle h",
                    "Non-work h", "Overtime h", "Tickets touched", "Emails sent"],
               [[b.user.name, b.status] + [day_fields(b)[k] for k in ("worked_hours", "productive_hours", "meeting_hours",
                                                                       "communication_hours", "idle_hours", "non_work_hours",
                                                                       "overtime_hours", "tickets_touched", "emails_sent")]
                for b in bundles], [22, 10] + [12] * 9)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --- company template ----------------------------------------------------------------------------------
def _cell_ref(ref: str) -> tuple[int, int]:
    col, row = coordinate_from_string(ref.strip().upper())
    return row, column_index_from_string(col)


def validate_mapping(mapping: dict, sheetnames: list[str]) -> list[str]:
    from app.exports.bundle import FIELD_HELP
    errors = []
    if mapping.get("sheet") and mapping["sheet"] not in sheetnames:
        errors.append(f"No sheet called {mapping['sheet']!r}; the workbook has {sheetnames}")
    for field, ref in (mapping.get("cells") or {}).items():
        if field not in FIELD_HELP:
            errors.append(f"Unknown field {field!r}")
        try:
            _cell_ref(ref)
        except (ValueError, TypeError, AttributeError, CellCoordinatesException):
            errors.append(f"{field}: {ref!r} is not a cell like B3")
    cols = mapping.get("columns") or {}
    for field, col in cols.items():
        if field not in FIELD_HELP:
            errors.append(f"Unknown field {field!r}")
        try:
            column_index_from_string(str(col).strip().upper())
        except ValueError:
            errors.append(f"{field}: {col!r} is not a column letter")
    if cols and not isinstance(mapping.get("start_row"), int):
        errors.append("start_row (the first task row) is required when columns are mapped")
    if not (mapping.get("cells") or cols):
        errors.append("Map at least one field to a cell or a column")
    return errors


def fill_template(template: bytes, mapping: dict, bundles: list[Bundle]) -> bytes:
    """Fixed cells get the first person-day's values; mapped columns get one row per task across all bundles.
    Styles, formulas and other sheets are kept (openpyxl does not keep charts or images in the template)."""
    wb = load_workbook(io.BytesIO(template))
    ws = wb[mapping["sheet"]] if mapping.get("sheet") else wb.active
    first = day_fields(bundles[0]) if bundles else {}
    for field, ref in (mapping.get("cells") or {}).items():
        r, c = _cell_ref(ref)
        ws.cell(row=r, column=c, value=first.get(field, ""))
    cols = {f: column_index_from_string(str(c).strip().upper()) for f, c in (mapping.get("columns") or {}).items()}
    if cols:
        start = mapping["start_row"]
        style_row = {c: ws.cell(row=start, column=c) for c in cols.values()}
        r = start
        for b in bundles:
            dfields = day_fields(b)
            for t in b.content.get("tasks", []) or [{"title": "", "seconds": 0, "status": "", "refs": []}]:
                values = {**dfields, **task_fields(t)}
                for field, c in cols.items():
                    cell = ws.cell(row=r, column=c, value=values.get(field, ""))
                    if r != start and style_row[c].has_style:
                        cell._style = copy(style_row[c]._style)
                r += 1
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def preview(filled: bytes, sheet: str | None, max_rows: int = 30, max_cols: int = 12) -> list[list]:
    ws = load_workbook(io.BytesIO(filled), data_only=False)
    ws = ws[sheet] if sheet else ws.active
    out = []
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, max_rows), max_col=min(ws.max_column, max_cols)):
        out.append(["" if c.value is None else str(c.value) for c in row])
    return out


# --- master workbook -----------------------------------------------------------------------------------
MASTER_HEADER = ["Date", "Person", "Email", "Team", "Task", "Summary", "Hours", "Status", "Category", "Tickets",
                 "Worked h", "Productive h", "Meetings h", "Tickets responded", "Tickets resolved", "Emails sent",
                 "Highlights", "Blockers", "Tomorrow", "AI insights"]


def append_master(path: Path, b: Bundle) -> int:
    """One sheet per month, one row per task. Re-running a person-day replaces that day's rows (no duplicates)."""
    wb = load_workbook(path) if path.exists() else Workbook()
    if not path.exists():
        wb.remove(wb.active)
    name = b.day.strftime("%Y-%m")
    if name not in wb.sheetnames:
        ws = wb.create_sheet(name)
        ws.append(MASTER_HEADER)
        for c in ws[1]:
            c.font, c.fill = HEAD, HEAD_FILL
        ws.freeze_panes = "A2"
    ws = wb[name]
    for r in range(ws.max_row, 1, -1):
        d, email = ws.cell(row=r, column=1).value, ws.cell(row=r, column=3).value
        d = d.date() if hasattr(d, "date") else d
        if d == b.day and email == b.user.email:
            ws.delete_rows(r)
    f = day_fields(b)
    n = 0
    for t in b.content.get("tasks", []):
        tf = task_fields(t)
        ws.append([b.day, b.user.name, b.user.email, b.team, tf["task"], tf["task_summary"], tf["hours"],
                   tf["task_status"], tf["category"], tf["tickets"], f["worked_hours"], f["productive_hours"],
                   f["meeting_hours"], f["tickets_responded"], f["tickets_resolved"], f["emails_sent"],
                   f["highlights"], f["blockers"], f["tomorrow"], f["ai_summary"]])
        n += 1
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return n
