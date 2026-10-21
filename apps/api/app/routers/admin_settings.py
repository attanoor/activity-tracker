"""Admin: settings and integrations, category rules and AI rule suggestions, the company Excel template."""

import json
import re
import secrets
import time as time_mod
from datetime import UTC, date, datetime, timedelta
from pathlib import Path

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from openpyxl import load_workbook
from pydantic import BaseModel, Field
from sqlalchemy import select
from sqlalchemy.orm import Session

from app import settings_store
from app.ai.llm import LLMUnavailable, get_llm
from app.audit import audit
from app.config import ROOT
from app.db import get_db
from app.exports import excel
from app.exports.bundle import FIELD_HELP, TASK_FIELD_NAMES, gather
from app.exports.mailer import MailNotConfigured, send
from app.integrations import mantis, runner
from app.models import Category, CategoryRule, ExcelTemplate, RuleSuggestion, User
from app.routers.admin import admin_only

router = APIRouter(prefix="/admin", tags=["admin"])
TEMPLATES = ROOT / ".local" / "templates"


# --- settings ------------------------------------------------------------------------------------------
@router.get("/settings")
def get_settings(db: Session = Depends(get_db), _: User = Depends(admin_only)):
    out = {k: settings_store.public(db, k) for k in settings_store.DEFAULTS}
    out["mantis"]["token_set"] = bool(mantis_token_present())
    out["mantis"]["token_env"] = mantis.TOKEN_ENV
    return out


def mantis_token_present() -> bool:
    try:
        return bool(mantis.token())
    except mantis.MantisNotConfigured:
        return False


@router.put("/settings/{section}")
def put_settings(section: str, body: dict, db: Session = Depends(get_db), admin: User = Depends(admin_only)):
    if section not in settings_store.DEFAULTS:
        raise HTTPException(404, "Unknown settings section")
    if section == "google" and body.get("service_account_json"):
        try:
            key = json.loads(body["service_account_json"])
            assert key.get("type") == "service_account" and key.get("private_key") and key.get("client_email")
        except (ValueError, AssertionError):
            raise HTTPException(422, "That is not a Google service-account key file (JSON)")
    if section == "google" and body.get("mode") not in (None, "off", "service_account", "oauth"):
        raise HTTPException(422, "mode must be off, service_account or oauth")
    out = settings_store.update(db, section, body)
    audit(db, admin.id, "settings.update", section, {"fields": sorted(k for k in body if "secret" not in k
                                                                         and "password" not in k and "json" not in k)})
    db.commit()
    return out


class TestEmail(BaseModel):
    to: str = Field(min_length=3)


@router.post("/settings/email/test")
def test_email(body: TestEmail, db: Session = Depends(get_db), _: User = Depends(admin_only)):
    try:
        send(db, [body.to], "Activity Tracker test email", "Email from Activity Tracker works.",
             "<p>Email from Activity Tracker works.</p>")
    except MailNotConfigured as e:
        raise HTTPException(400, str(e))
    except Exception as e:  # SMTP and network errors are reported to the admin as-is
        raise HTTPException(502, f"Sending failed: {e}")
    return {"ok": True}


@router.post("/settings/ai/test")
def test_ai(db: Session = Depends(get_db), _: User = Depends(admin_only)):
    llm, model = get_llm(db)
    if not llm:
        raise HTTPException(400, "AI is switched off")
    t = time_mod.monotonic()
    try:
        raw = llm("Reply with JSON only.", 'Reply with {"ok": true}')
    except LLMUnavailable as e:
        raise HTTPException(502, f"Could not reach the model: {e}")
    return {"model": model, "seconds": round(time_mod.monotonic() - t, 1), "reply": raw[:200]}


class SyncIn(BaseModel):
    user_id: int | None = None
    day: date
    sources: list[str] = ["google", "mantis"]


@router.post("/sync")
def sync_now(body: SyncIn, db: Session = Depends(get_db), _: User = Depends(admin_only)):
    out: dict = {}
    if "mantis" in body.sources:
        out["mantis"] = runner.sync_mantis(db, datetime.combine(body.day, datetime.min.time(), UTC) - timedelta(hours=14))
    if "google" in body.sources:
        users = [db.get(User, body.user_id)] if body.user_id else list(db.scalars(select(User).where(User.active.is_(True))))
        out["google"] = {u.name: runner.sync_google(db, u, body.day) for u in users if u}
    return out


# --- category rules ------------------------------------------------------------------------------------
@router.get("/categories")
def categories(db: Session = Depends(get_db), _: User = Depends(admin_only)):
    return [{"id": c.id, "name": c.name, "klass": c.klass} for c in db.scalars(select(Category).order_by(Category.name))]


class RuleIn(BaseModel):
    field: str = Field(pattern="^(domain|app|title)$")
    pattern: str = Field(min_length=1, max_length=255)
    category_id: int
    priority: int = 0
    enabled: bool = True


def _rule_out(r: CategoryRule) -> dict:
    return {"id": r.id, "field": r.field, "pattern": r.pattern, "category_id": r.category_id, "priority": r.priority,
            "enabled": r.enabled}


def _check_rule(db: Session, body: RuleIn) -> None:
    if not db.get(Category, body.category_id):
        raise HTTPException(422, "Unknown category")
    if body.field == "title":
        try:
            re.compile(body.pattern)
        except re.error as e:
            raise HTTPException(422, f"Not a valid pattern: {e}")


@router.get("/rules")
def rules(db: Session = Depends(get_db), _: User = Depends(admin_only)):
    return [_rule_out(r) for r in db.scalars(select(CategoryRule).order_by(CategoryRule.priority.desc(), CategoryRule.field,
                                                                          CategoryRule.pattern))]


@router.post("/rules", status_code=201)
def add_rule(body: RuleIn, db: Session = Depends(get_db), admin: User = Depends(admin_only)):
    _check_rule(db, body)
    r = CategoryRule(**body.model_dump())
    db.add(r)
    db.flush()
    audit(db, admin.id, "rule.create", f"rule:{r.id}", body.model_dump())
    db.commit()
    return _rule_out(r)


@router.put("/rules/{rule_id}")
def edit_rule(rule_id: int, body: RuleIn, db: Session = Depends(get_db), admin: User = Depends(admin_only)):
    r = db.get(CategoryRule, rule_id)
    if not r:
        raise HTTPException(404, "Unknown rule")
    _check_rule(db, body)
    for k, v in body.model_dump().items():
        setattr(r, k, v)
    audit(db, admin.id, "rule.update", f"rule:{r.id}", body.model_dump())
    db.commit()
    return _rule_out(r)


@router.delete("/rules/{rule_id}", status_code=204)
def delete_rule(rule_id: int, db: Session = Depends(get_db), admin: User = Depends(admin_only)):
    r = db.get(CategoryRule, rule_id)
    if r:
        audit(db, admin.id, "rule.delete", f"rule:{r.id}", _rule_out(r))
        db.delete(r)
        db.commit()


@router.get("/rule-suggestions")
def rule_suggestions(db: Session = Depends(get_db), _: User = Depends(admin_only)):
    return [{"id": s.id, "domain": s.domain, "category_id": s.category_id, "reason": s.reason,
             "seconds_seen": s.seconds_seen, "status": s.status}
            for s in db.scalars(select(RuleSuggestion).where(RuleSuggestion.status == "pending")
                                .order_by(RuleSuggestion.seconds_seen.desc()))]


class Decision(BaseModel):
    accept: bool
    category_id: int | None = None  # the admin may pick a different category when accepting


@router.post("/rule-suggestions/{sid}")
def decide(sid: int, body: Decision, db: Session = Depends(get_db), admin: User = Depends(admin_only)):
    s = db.get(RuleSuggestion, sid)
    if not s or s.status != "pending":
        raise HTTPException(404, "Unknown or already decided")
    s.status = "accepted" if body.accept else "rejected"
    if body.accept:
        cat = body.category_id or s.category_id
        if not db.get(Category, cat):
            raise HTTPException(422, "Unknown category")
        db.add(CategoryRule(field="domain", pattern=s.domain, category_id=cat))
    audit(db, admin.id, "rule_suggestion." + s.status, s.domain)
    db.commit()
    return {"id": s.id, "status": s.status}


# --- company Excel template ----------------------------------------------------------------------------
def template_path(t: ExcelTemplate) -> Path:
    return TEMPLATES / t.file_name


def _tpl_out(t: ExcelTemplate) -> dict:
    sheets = []
    try:
        sheets = load_workbook(template_path(t), read_only=True).sheetnames
    except Exception:  # a missing or unreadable file shows as no sheets
        pass
    return {"id": t.id, "name": t.name, "mapping": t.mapping, "active": t.active, "sheets": sheets,
            "created_at": t.created_at}


@router.get("/templates")
def templates(db: Session = Depends(get_db), _: User = Depends(admin_only)):
    return {"templates": [_tpl_out(t) for t in db.scalars(select(ExcelTemplate).order_by(ExcelTemplate.id.desc()))],
            "fields": FIELD_HELP, "task_fields": sorted(TASK_FIELD_NAMES)}


@router.post("/templates", status_code=201)
async def upload_template(name: str = Form(...), file: UploadFile = File(...), db: Session = Depends(get_db),
                          admin: User = Depends(admin_only)):
    data = await file.read()
    if len(data) > 10 * 1024 * 1024:
        raise HTTPException(413, "Template is larger than 10 MB")
    import io
    try:
        load_workbook(io.BytesIO(data))
    except Exception:
        raise HTTPException(422, "That file is not an Excel workbook (.xlsx)")
    TEMPLATES.mkdir(parents=True, exist_ok=True)
    fname = f"{secrets.token_hex(8)}.xlsx"
    (TEMPLATES / fname).write_bytes(data)
    t = ExcelTemplate(name=name[:120], file_name=fname, mapping={})
    db.add(t)
    db.flush()
    audit(db, admin.id, "template.upload", f"template:{t.id}", {"name": t.name})
    db.commit()
    return _tpl_out(t)


@router.put("/templates/{tid}/mapping")
def set_mapping(tid: int, mapping: dict, db: Session = Depends(get_db), admin: User = Depends(admin_only)):
    t = db.get(ExcelTemplate, tid)
    if not t:
        raise HTTPException(404, "Unknown template")
    errors = excel.validate_mapping(mapping, _tpl_out(t)["sheets"])
    if errors:
        raise HTTPException(422, "; ".join(errors))
    t.mapping = mapping
    audit(db, admin.id, "template.mapping", f"template:{t.id}")
    db.commit()
    return _tpl_out(t)


class PreviewIn(BaseModel):
    user_id: int
    day: date
    mapping: dict | None = None


@router.post("/templates/{tid}/preview")
def preview(tid: int, body: PreviewIn, db: Session = Depends(get_db), _: User = Depends(admin_only)):
    t = db.get(ExcelTemplate, tid)
    u = db.get(User, body.user_id)
    if not t or not u:
        raise HTTPException(404, "Unknown template or person")
    mapping = body.mapping or t.mapping
    errors = excel.validate_mapping(mapping, _tpl_out(t)["sheets"])
    if errors:
        raise HTTPException(422, "; ".join(errors))
    filled = excel.fill_template(template_path(t).read_bytes(), mapping, [gather(db, u, body.day)])
    return {"rows": excel.preview(filled, mapping.get("sheet"))}


@router.post("/templates/{tid}/activate")
def activate(tid: int, db: Session = Depends(get_db), admin: User = Depends(admin_only)):
    t = db.get(ExcelTemplate, tid)
    if not t:
        raise HTTPException(404, "Unknown template")
    if excel.validate_mapping(t.mapping, _tpl_out(t)["sheets"]):
        raise HTTPException(422, "Save a valid mapping first")
    for other in db.scalars(select(ExcelTemplate)):
        other.active = other.id == t.id
    audit(db, admin.id, "template.activate", f"template:{t.id}")
    db.commit()
    return _tpl_out(t)


@router.delete("/templates/{tid}", status_code=204)
def delete_template(tid: int, db: Session = Depends(get_db), admin: User = Depends(admin_only)):
    t = db.get(ExcelTemplate, tid)
    if t:
        template_path(t).unlink(missing_ok=True)
        audit(db, admin.id, "template.delete", f"template:{t.id}")
        db.delete(t)
        db.commit()
